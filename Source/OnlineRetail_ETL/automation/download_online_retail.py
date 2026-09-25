#!/usr/bin/env python3
from __future__ import annotations
import argparse,csv,hashlib,json,sys,urllib.request,zipfile
from datetime import datetime
from pathlib import Path
PRIMARY_URL='https://archive.ics.uci.edu/static/public/352/online%2Bretail.zip'
DATASET_PAGE='https://archive.ics.uci.edu/dataset/352/online%2Bretail'
DOI='10.24432/C5BW33'; EXPECTED_ROWS=541_909
EXPECTED_HEADER=['InvoiceNo','StockCode','Description','Quantity','InvoiceDate','UnitPrice','CustomerID','Country']
UA='Mozilla/5.0 (X11; Linux x86_64) AppleWebKit/537.36 BTA12-UIT/1.0'
def sha256(p):
 h=hashlib.sha256()
 with open(p,'rb') as f:
  for b in iter(lambda:f.read(1024*1024),b''): h.update(b)
 return h.hexdigest()
def download_limited(url,dest,cap):
 dest=Path(dest); tmp=dest.with_suffix(dest.suffix+'.partial'); tmp.unlink(missing_ok=True); dest.parent.mkdir(parents=True,exist_ok=True)
 req=urllib.request.Request(url,headers={'User-Agent':UA,'Accept':'application/zip,application/octet-stream,*/*;q=0.8','Referer':'https://archive.ics.uci.edu/'})
 try:
  with urllib.request.urlopen(req,timeout=120) as r:
   cl=r.headers.get('Content-Length')
   if cl and int(cl)>cap: raise RuntimeError(f'Remote file is {int(cl):,} bytes > {cap:,}')
   n=0
   with tmp.open('wb') as f:
    while True:
     b=r.read(1024*1024)
     if not b: break
     n+=len(b)
     if n>cap: raise RuntimeError(f'Download crossed {cap:,} bytes')
     f.write(b); print(f'[INFO] download {n/1_000_000:.2f} MB / {cap/1_000_000:.2f} MB',flush=True)
 except Exception:
  tmp.unlink(missing_ok=True); raise
 tmp.replace(dest); return {'bytes':dest.stat().st_size,'sha256':sha256(dest)}
def extract_xlsx(zpath,raw,cap):
 with zipfile.ZipFile(zpath) as z:
  xs=[x for x in z.infolist() if x.filename.lower().endswith('.xlsx') and not x.is_dir()]
  if len(xs)!=1: raise RuntimeError(f'Expected one XLSX, found {[x.filename for x in xs]}')
  x=xs[0]
  if x.file_size>cap: raise RuntimeError('XLSX exceeds 50MB hard cap')
  out=Path(raw)/'Online Retail.xlsx'; tmp=out.with_suffix('.xlsx.partial')
  with z.open(x) as s,tmp.open('wb') as d:
   n=0
   while True:
    b=s.read(1024*1024)
    if not b: break
    n+=len(b)
    if n>cap: raise RuntimeError('Extracted XLSX exceeds hard cap')
    d.write(b)
  tmp.replace(out); return out
def norm(v,col):
 if v is None:return ''
 if col=='InvoiceDate' and isinstance(v,datetime): return v.strftime('%Y-%m-%d %H:%M:%S')
 if col in ('InvoiceNo','StockCode','CustomerID','Quantity') and isinstance(v,float) and v.is_integer(): return str(int(v))
 if col in ('InvoiceNo','StockCode','CustomerID','Quantity') and isinstance(v,int): return str(v)
 if col=='UnitPrice' and isinstance(v,(int,float)): return format(v,'.10g')
 return str(v).strip()
def convert(xlsx,csvpath,cap):
 from openpyxl import load_workbook
 wb=load_workbook(xlsx,read_only=True,data_only=True); ws=wb[wb.sheetnames[0]]; it=ws.iter_rows(values_only=True)
 h=[str(x).strip() if x is not None else '' for x in next(it)]
 if h!=EXPECTED_HEADER: raise RuntimeError(f'Unexpected header: {h!r}')
 csvpath=Path(csvpath); tmp=csvpath.with_suffix(csvpath.suffix+'.partial'); tmp.unlink(missing_ok=True); csvpath.parent.mkdir(parents=True,exist_ok=True)
 n=0
 try:
  with tmp.open('w',encoding='utf-8',newline='') as f:
   w=csv.writer(f,lineterminator='\n'); w.writerow(EXPECTED_HEADER)
   for row in it:
    n+=1; w.writerow([norm(v,EXPECTED_HEADER[i]) for i,v in enumerate(row)])
    if n%10000==0:
     f.flush(); size=tmp.stat().st_size; print(f'[INFO] converted {n:,}/{EXPECTED_ROWS:,}; CSV={size/1_000_000:.2f} MB',flush=True)
     if size>cap: raise RuntimeError(f'Generated CSV > {cap:,} bytes')
 except Exception:
  tmp.unlink(missing_ok=True); raise
 wb.close()
 if n!=EXPECTED_ROWS: tmp.unlink(missing_ok=True); raise RuntimeError(f'Expected {EXPECTED_ROWS:,} rows, got {n:,}')
 if tmp.stat().st_size>cap: tmp.unlink(missing_ok=True); raise RuntimeError('Generated CSV exceeds cap')
 tmp.replace(csvpath); return n
def verify(p):
 with open(p,encoding='utf-8',newline='') as f:
  r=csv.reader(f); h=next(r)
  if h!=EXPECTED_HEADER: raise RuntimeError('CSV header mismatch')
  return sum(1 for _ in r)
def main():
 ap=argparse.ArgumentParser(); ap.add_argument('--output',required=True);ap.add_argument('--manifest',required=True);ap.add_argument('--raw-dir',required=True);ap.add_argument('--max-bytes',type=int,default=50_000_000);a=ap.parse_args()
 if a.max_bytes>50_000_000: raise RuntimeError('Refusing cap >50,000,000 bytes')
 raw=Path(a.raw_dir).resolve(); raw.mkdir(parents=True,exist_ok=True); z=raw/'online_retail_uci.zip'
 print('[INFO] Downloading COMPLETE UCI Online Retail dataset (no sampling).')
 zm=download_limited(PRIMARY_URL,z,a.max_bytes); x=extract_xlsx(z,raw,a.max_bytes); xm={'bytes':x.stat().st_size,'sha256':sha256(x)}
 print(f"[PASS] source XLSX={xm['bytes']:,} bytes")
 n=convert(x,Path(a.output),a.max_bytes); v=verify(a.output)
 if n!=EXPECTED_ROWS or v!=EXPECTED_ROWS: raise RuntimeError('Final row verification failed')
 out=Path(a.output).resolve(); m={'dataset':'UCI Online Retail','uci_dataset_id':352,'doi':DOI,'dataset_page':DATASET_PAGE,'download_url':PRIMARY_URL,'scope':'complete dataset','sampling':False,'truncation':False,'expected_rows':EXPECTED_ROWS,'rows_converted':n,'rows_verified':v,'max_bytes':a.max_bytes,'archive':{'path':str(z),'bytes':zm['bytes'],'sha256':zm['sha256']},'xlsx':{'path':str(x),'bytes':xm['bytes'],'sha256':xm['sha256']},'csv':{'path':str(out),'bytes':out.stat().st_size,'sha256':sha256(out)},'columns':EXPECTED_HEADER,'generated_at_utc':datetime.utcnow().replace(microsecond=0).isoformat()+'Z'}
 Path(a.manifest).write_text(json.dumps(m,indent=2),encoding='utf-8'); print(f"[PASS] COMPLETE dataset rows={v:,}, CSV={out.stat().st_size:,} bytes")
if __name__=='__main__':
 try: main()
 except Exception as e: print(f'[FAIL] {e}',file=sys.stderr); raise
